"""Tests for the Gemini parser (IMP-011) — CSV and native xlsx paths."""

import tempfile
import os
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from imports.exchanges.gemini import GeminiImporter, _parse_number, _to_float, _fmt_dt
from imports.registry import register, get_parser, get_all_parsers, clear_registry
from imports.validation import validate_batch
from db import SqliteBackend
from cryptoAccounts import CryptoAccounts

FIXTURES_DIR = Path(__file__).parent / "fixtures" / "csv_samples"

# Shared header — hyphenated column names match Gemini's export
_HDR = "time,base-asset,quote-asset,type,price,quantity,total,fee,fee-currency,trade-id\n"


def _make_csv(content: str) -> str:
    """Write content to a temp CSV file and return its path."""
    f = tempfile.NamedTemporaryFile(mode='w', suffix='.csv', delete=False)
    f.write(content)
    f.close()
    return f.name


class TestGeminiImporterDetection:
    """Detection logic for Gemini CSV exports."""

    def test_detects_gemini_export(self):
        """File with base-asset, quote-asset, trade-id is detected."""
        csv_path = _make_csv(
            _HDR +
            "2024-01-05 10:00:00,BTC,USD,Buy,29500,0.01,295,2.95,USD,t-1\n"
        )
        try:
            parser = GeminiImporter()
            assert parser.detect(csv_path) is True
        finally:
            os.unlink(csv_path)

    def test_rejects_cashapp_export(self):
        """Cash App CSV (no hyphenated columns) is not detected as Gemini."""
        csv_path = _make_csv(
            "Date,Transaction Type,Amount (BTC),Market Price ($),"
            "Cost Basis ($),Proceeds ($),Gain/Loss ($),Note\n"
            "2024-01-10 10:30:00,Purchase,0.01,29500.00,295.00,,,\n"
        )
        try:
            parser = GeminiImporter()
            assert parser.detect(csv_path) is False
        finally:
            os.unlink(csv_path)

    def test_rejects_coinbase_export(self):
        """Coinbase CSV lacks Gemini-specific hyphenated columns."""
        csv_path = _make_csv(
            "Timestamp,Transaction Type,Asset,Quantity Transacted,Spot Price Currency,"
            "Spot Price at Transaction,Subtotal,Total,Fees,Notes\n"
            "2024-01-15 10:00:00 UTC,Buy,BTC,0.05,USD,30000,1500,1502.50,2.50,\n"
        )
        try:
            parser = GeminiImporter()
            assert parser.detect(csv_path) is False
        finally:
            os.unlink(csv_path)

    def test_rejects_empty_file(self):
        """Empty file is not detected."""
        csv_path = _make_csv("")
        try:
            parser = GeminiImporter()
            assert parser.detect(csv_path) is False
        finally:
            os.unlink(csv_path)


class TestGeminiParseNumber:
    """Unit tests for the _parse_number helper."""

    def test_plain_number(self):
        assert _parse_number("0.01") == 0.01

    def test_dollar_sign_and_commas(self):
        assert _parse_number("$29,500.00") == 29500.0

    def test_empty_string(self):
        assert _parse_number("") == 0.0

    def test_whitespace_only(self):
        assert _parse_number("   ") == 0.0

    def test_non_numeric(self):
        assert _parse_number("abc") == 0.0

    def test_dollar_sign_only(self):
        assert _parse_number("$") == 0.0


class TestGeminiImporterParsing:
    """Row-level parsing for each Gemini transaction type."""

    def test_parse_fixture_file(self):
        """Fixture: 1 Buy + 1 Sell = 2 Trade, 1 Deposit, 1 Withdrawal, 1 Earn; ETH row filtered."""
        parser = GeminiImporter()
        colnames, transactions = parser.parse(str(FIXTURES_DIR / "gemini_sample.csv"))

        assert len(colnames) == 11
        # 6 rows - 1 ETH = 5 transactions
        assert len(transactions) == 5

        type_counts: dict[str, int] = {}
        for tx in transactions:
            type_counts[tx['trans_type']] = type_counts.get(tx['trans_type'], 0) + 1
        assert type_counts.get('Trade', 0) == 2
        assert type_counts.get('Deposit', 0) == 1
        assert type_counts.get('Withdrawal', 0) == 1
        assert type_counts.get('Interest Income', 0) == 1

    def test_buy_maps_to_trade(self):
        """Buy row: BTC is bought, USD (quote-asset) is sold; fee in USD."""
        csv_path = _make_csv(
            _HDR +
            "2024-01-05 10:00:00,BTC,USD,Buy,29500.00,0.01,295.00,2.95,USD,trade-001\n"
        )
        try:
            parser = GeminiImporter()
            _, txs = parser.parse(csv_path)
            assert len(txs) == 1
            tx = txs[0]

            assert tx['trans_type'] == 'Trade'
            assert tx['exchange'] == 'Gemini'
            assert tx['buy'] == 0.01
            assert tx['buy_curr'] == 'BTC'
            assert tx['sell'] == 295.00      # total
            assert tx['sell_curr'] == 'USD'
            assert tx['fee'] == 2.95
            assert tx['fee_curr'] == 'USD'
        finally:
            os.unlink(csv_path)

    def test_buy_falls_back_to_price_times_quantity(self):
        """Buy with empty total computes sell from price * quantity."""
        csv_path = _make_csv(
            _HDR +
            "2024-01-05 10:00:00,BTC,USD,Buy,29500.00,0.01,,0,USD,t-1\n"
        )
        try:
            parser = GeminiImporter()
            _, txs = parser.parse(csv_path)
            assert len(txs) == 1
            # 29500 * 0.01 = 295.0
            assert txs[0]['sell'] == pytest.approx(295.0)
        finally:
            os.unlink(csv_path)

    def test_sell_maps_to_trade(self):
        """Sell row: BTC is sold, USD (proceeds) is bought."""
        csv_path = _make_csv(
            _HDR +
            "2024-01-20 14:30:00,BTC,USD,Sell,30000.00,0.005,150.00,1.50,USD,trade-002\n"
        )
        try:
            parser = GeminiImporter()
            _, txs = parser.parse(csv_path)
            assert len(txs) == 1
            tx = txs[0]

            assert tx['trans_type'] == 'Trade'
            assert tx['buy'] == 150.00       # proceeds (total)
            assert tx['buy_curr'] == 'USD'
            assert tx['sell'] == 0.005       # BTC sold
            assert tx['sell_curr'] == 'BTC'
            assert tx['fee'] == 1.50
            assert tx['fee_curr'] == 'USD'
        finally:
            os.unlink(csv_path)

    def test_sell_falls_back_to_price_times_quantity(self):
        """Sell with empty total computes buy from price * quantity."""
        csv_path = _make_csv(
            _HDR +
            "2024-01-20 14:30:00,BTC,USD,Sell,30000.00,0.005,,0,USD,t-2\n"
        )
        try:
            parser = GeminiImporter()
            _, txs = parser.parse(csv_path)
            assert len(txs) == 1
            # 30000 * 0.005 = 150.0
            assert txs[0]['buy'] == pytest.approx(150.0)
        finally:
            os.unlink(csv_path)

    def test_deposit_maps_correctly(self):
        """Deposit row: BTC received, zero fee when fee field is empty."""
        csv_path = _make_csv(
            _HDR +
            "2024-02-01 09:00:00,BTC,USD,Deposit,,0.02,,,,\n"
        )
        try:
            parser = GeminiImporter()
            _, txs = parser.parse(csv_path)
            assert len(txs) == 1
            tx = txs[0]

            assert tx['trans_type'] == 'Deposit'
            assert tx['exchange'] == 'Gemini'
            assert tx['buy'] == 0.02
            assert tx['buy_curr'] == 'BTC'
            assert tx['sell'] == 0.0
            assert tx['fee'] == 0.0
            assert tx['fee_curr'] == ''
        finally:
            os.unlink(csv_path)

    def test_withdrawal_default_exchange(self):
        """Withdrawal with withdraw_to=None uses 'Gemini-Withdrawal' + review comment."""
        csv_path = _make_csv(
            _HDR +
            "2024-02-15 16:00:00,BTC,USD,Withdrawal,,0.01,,0.0001,BTC,\n"
        )
        try:
            parser = GeminiImporter()
            _, txs = parser.parse(csv_path, withdraw_to=None)
            assert len(txs) == 1
            tx = txs[0]

            assert tx['trans_type'] == 'Withdrawal'
            assert tx['exchange'] == 'Gemini-Withdrawal'
            assert tx['sell'] == 0.01
            assert tx['sell_curr'] == 'BTC'
            assert 'Review' in tx['comment']
        finally:
            os.unlink(csv_path)

    def test_withdrawal_with_withdraw_to(self):
        """Withdrawal with explicit withdraw_to sets exchange; no review comment."""
        csv_path = _make_csv(
            _HDR +
            "2024-02-15 16:00:00,BTC,USD,Withdrawal,,0.01,,0.0001,BTC,\n"
        )
        try:
            parser = GeminiImporter()
            _, txs = parser.parse(csv_path, withdraw_to="ColdCard")
            assert len(txs) == 1
            tx = txs[0]

            assert tx['exchange'] == 'ColdCard'
            assert 'Review' not in tx['comment']
        finally:
            os.unlink(csv_path)

    def test_withdrawal_btc_fee_tracked(self):
        """Withdrawal with a BTC network fee is recorded correctly."""
        csv_path = _make_csv(
            _HDR +
            "2024-02-15 16:00:00,BTC,USD,Withdrawal,,0.01,,0.0001,BTC,\n"
        )
        try:
            parser = GeminiImporter()
            _, txs = parser.parse(csv_path, withdraw_to="Ledger")
            assert txs[0]['fee'] == 0.0001
            assert txs[0]['fee_curr'] == 'BTC'
        finally:
            os.unlink(csv_path)

    def test_earn_maps_to_interest_income(self):
        """Earn row maps to Interest Income with 'Gemini Earn' comment."""
        csv_path = _make_csv(
            _HDR +
            "2024-02-20 12:00:00,BTC,USD,Earn,,0.001,,,,trade-003\n"
        )
        try:
            parser = GeminiImporter()
            _, txs = parser.parse(csv_path)
            assert len(txs) == 1
            tx = txs[0]

            assert tx['trans_type'] == 'Interest Income'
            assert tx['exchange'] == 'Gemini'
            assert tx['buy'] == 0.001
            assert tx['buy_curr'] == 'BTC'
            assert tx['sell'] == 0.0
            assert tx['comment'] == 'Gemini Earn'
        finally:
            os.unlink(csv_path)

    def test_earn_interest_alias(self):
        """'Earn Interest' type (alternate label) also maps to Interest Income."""
        csv_path = _make_csv(
            _HDR +
            "2024-02-20 12:00:00,BTC,USD,Earn Interest,,0.0005,,,,\n"
        )
        try:
            parser = GeminiImporter()
            _, txs = parser.parse(csv_path)
            assert len(txs) == 1
            assert txs[0]['trans_type'] == 'Interest Income'
        finally:
            os.unlink(csv_path)

    def test_non_btc_filtered_out(self):
        """ETH and other non-BTC base-assets are silently dropped."""
        csv_path = _make_csv(
            _HDR +
            "2024-01-15 11:00:00,ETH,USD,Buy,2200.00,1.5,3300.00,3.30,USD,t-4\n"
            "2024-01-16 11:00:00,SOL,USD,Buy,100.00,10,1000.00,1.00,USD,t-5\n"
        )
        try:
            parser = GeminiImporter()
            _, txs = parser.parse(csv_path)
            assert len(txs) == 0
        finally:
            os.unlink(csv_path)

    def test_row_missing_time_is_skipped(self):
        """Row where time is empty is silently dropped."""
        csv_path = _make_csv(
            _HDR +
            ",BTC,USD,Buy,29500.00,0.01,295.00,2.95,USD,t-1\n"
        )
        try:
            parser = GeminiImporter()
            _, txs = parser.parse(csv_path)
            assert len(txs) == 0
        finally:
            os.unlink(csv_path)

    def test_empty_file_returns_no_transactions(self):
        """File with only a header row yields an empty transaction list."""
        csv_path = _make_csv(_HDR)
        try:
            parser = GeminiImporter()
            colnames, txs = parser.parse(csv_path)
            assert len(txs) == 0
            assert len(colnames) == 11
        finally:
            os.unlink(csv_path)


class TestGeminiImporterRegistration:
    """Gemini parser registration and discovery."""

    def setup_method(self):
        clear_registry()
        register(GeminiImporter)

    def teardown_method(self):
        clear_registry()

    def test_registered_by_name(self):
        assert get_parser("Gemini") is not None

    def test_registered_parser_is_gemini_importer(self):
        assert isinstance(get_parser("Gemini"), GeminiImporter)

    def test_appears_in_all_parsers(self):
        names = [p.name for p in get_all_parsers()]
        assert "Gemini" in names


class TestGeminiImporterIntegration:
    """End-to-end import from Gemini CSV into in-memory SQLite."""

    def setup_method(self):
        clear_registry()
        register(GeminiImporter)

    def teardown_method(self):
        clear_registry()

    def _make_backend(self):
        return SqliteBackend(':memory:', auto_create_tables=True)

    def test_full_import_workflow(self):
        """Parse fixture, import all; verify counts and types."""
        parser = GeminiImporter()
        _, transactions = parser.parse(
            str(FIXTURES_DIR / "gemini_sample.csv"), withdraw_to="Ledger"
        )

        result = validate_batch(transactions)
        assert result.is_valid

        backend = self._make_backend()
        crypto = CryptoAccounts(backend=backend)
        import_result = crypto.import_transactions(transactions)

        assert import_result['imported'] == 5
        assert import_result['skipped'] == 0

        rows = backend.execute("SELECT trans_type, COUNT(*) as cnt FROM ledger GROUP BY trans_type")
        type_counts = {row['trans_type']: row['cnt'] for row in rows}

        assert type_counts.get('Trade', 0) == 2
        assert type_counts.get('Deposit', 0) == 1
        assert type_counts.get('Withdrawal', 0) == 1
        assert type_counts.get('Interest Income', 0) == 1

    def test_earn_comment_persists(self):
        """Earn transactions retain 'Gemini Earn' comment after import."""
        parser = GeminiImporter()
        _, transactions = parser.parse(
            str(FIXTURES_DIR / "gemini_sample.csv"), withdraw_to="Ledger"
        )

        backend = self._make_backend()
        crypto = CryptoAccounts(backend=backend)
        crypto.import_transactions(transactions)

        rows = backend.execute(
            "SELECT comment FROM ledger WHERE trans_type = 'Interest Income'"
        )
        assert len(rows) == 1
        assert rows[0]['comment'] == 'Gemini Earn'


# ---------------------------------------------------------------------------
# xlsx helpers
# ---------------------------------------------------------------------------

# Minimal column set that the xlsx parser actually reads.  Extra columns
# (BAT, GUSD, …) are deliberately omitted — dynamic header mapping means
# the parser never assumes they exist.
_XLSX_HEADERS = [
    "Date", "Type", "Symbol", "Specification",
    "USD Amount USD", "Fee (USD) USD",
    "BTC Amount BTC", "Fee (BTC) BTC",
    "Withdrawal Destination",
]


def _make_xlsx(rows: list[list], headers: list[str] | None = None) -> str:
    """Write an xlsx workbook with *headers* (row 1) and *rows* (row 2+).

    Returns the path to the temp file.  Each element in *rows* is a list
    aligned positionally with *headers*.
    """
    if headers is None:
        headers = _XLSX_HEADERS
    wb = openpyxl.Workbook()
    ws = wb.active
    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=h)
    for row_idx, row_data in enumerate(rows, start=2):
        for col_idx, val in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    f = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
    wb.save(f.name)
    wb.close()
    f.close()
    return f.name


def _xlsx_row(
    date=None, typ="", symbol="", spec="",
    usd_amt=0.0, fee_usd=0.0, btc_amt=0.0, fee_btc=0.0, wd_dest="",
) -> list:
    """Build a single xlsx data row aligned to _XLSX_HEADERS."""
    return [date, typ, symbol, spec, usd_amt, fee_usd, btc_amt, fee_btc, wd_dest]


# ---------------------------------------------------------------------------
# xlsx helper unit tests
# ---------------------------------------------------------------------------

class TestGeminiXlsxHelpers:
    """Unit tests for module-level xlsx helpers."""

    def test_to_float_none(self):
        assert _to_float(None) == 0.0

    def test_to_float_int(self):
        assert _to_float(42) == 42.0

    def test_to_float_negative(self):
        assert _to_float(-1234.56) == -1234.56

    def test_to_float_string(self):
        assert _to_float("99.9") == 99.9

    def test_to_float_bad_string(self):
        assert _to_float("nope") == 0.0

    def test_fmt_dt_datetime(self):
        dt = datetime(2024, 3, 15, 9, 30, 0)
        assert _fmt_dt(dt) == "2024-03-15 09:30:00"

    def test_fmt_dt_string_passthrough(self):
        assert _fmt_dt("2024-03-15") == "2024-03-15"

    def test_fmt_dt_none(self):
        assert _fmt_dt(None) == ''


# ---------------------------------------------------------------------------
# xlsx detection
# ---------------------------------------------------------------------------

class TestGeminiXlsxDetection:
    """Detection of native Gemini xlsx exports."""

    def test_detects_gemini_xlsx(self):
        """xlsx with 'BTC Amount BTC' + 'Withdrawal Destination' is detected."""
        path = _make_xlsx([])  # header-only is fine for detection
        try:
            assert GeminiImporter().detect(path) is True
        finally:
            os.unlink(path)

    def test_rejects_xlsx_missing_btc_column(self):
        """xlsx without 'BTC Amount BTC' header is rejected."""
        headers = ["Date", "Type", "Symbol", "Withdrawal Destination"]
        path = _make_xlsx([], headers=headers)
        try:
            assert GeminiImporter().detect(path) is False
        finally:
            os.unlink(path)

    def test_rejects_xlsx_missing_wd_column(self):
        """xlsx without 'Withdrawal Destination' header is rejected."""
        headers = ["Date", "Type", "Symbol", "BTC Amount BTC"]
        path = _make_xlsx([], headers=headers)
        try:
            assert GeminiImporter().detect(path) is False
        finally:
            os.unlink(path)

    def test_csv_detection_still_works_alongside_xlsx(self):
        """CSV detection is unaffected by xlsx support being present."""
        csv_path = _make_csv(
            _HDR +
            "2024-01-05 10:00:00,BTC,USD,Buy,29500,0.01,295,2.95,USD,t-1\n"
        )
        try:
            assert GeminiImporter().detect(csv_path) is True
        finally:
            os.unlink(csv_path)

    def test_rejects_nonexistent_xlsx(self):
        """Non-existent .xlsx path does not crash — returns False."""
        assert GeminiImporter().detect("/tmp/_no_such_file_gemini.xlsx") is False


# ---------------------------------------------------------------------------
# xlsx parsing — row-level
# ---------------------------------------------------------------------------

class TestGeminiXlsxParsing:
    """Row-level parsing for each xlsx transaction type."""

    def test_buy_maps_to_trade(self):
        """BUY + BTCUSD → Trade; USD amounts are negative in real exports."""
        dt = datetime(2024, 1, 10, 12, 0, 0)
        path = _make_xlsx([
            _xlsx_row(date=dt, typ="Buy", symbol="BTCUSD",
                      usd_amt=-500.0, fee_usd=-5.0, btc_amt=0.016),
        ])
        try:
            _, txs = GeminiImporter().parse(path)
            assert len(txs) == 1
            tx = txs[0]
            assert tx['trans_type'] == 'Trade'
            assert tx['exchange'] == 'Gemini'
            assert tx['buy'] == pytest.approx(0.016)
            assert tx['buy_curr'] == 'BTC'
            assert tx['sell'] == pytest.approx(500.0)   # abs of -500
            assert tx['sell_curr'] == 'USD'
            assert tx['fee'] == pytest.approx(5.0)      # abs of -5
            assert tx['fee_curr'] == 'USD'
            assert tx['created_date'] == '2024-01-10 12:00:00'
        finally:
            os.unlink(path)

    def test_buy_with_positive_amounts(self):
        """BUY works correctly even when amounts are already positive."""
        dt = datetime(2024, 2, 1, 8, 0, 0)
        path = _make_xlsx([
            _xlsx_row(date=dt, typ="Buy", symbol="BTCUSD",
                      usd_amt=1000.0, fee_usd=10.0, btc_amt=0.033),
        ])
        try:
            _, txs = GeminiImporter().parse(path)
            assert txs[0]['sell'] == pytest.approx(1000.0)
            assert txs[0]['fee'] == pytest.approx(10.0)
        finally:
            os.unlink(path)

    def test_sell_maps_to_trade(self):
        """SELL + BTCUSD → Trade; BTC amount negative, USD positive."""
        dt = datetime(2024, 3, 5, 15, 30, 0)
        path = _make_xlsx([
            _xlsx_row(date=dt, typ="Sell", symbol="BTCUSD",
                      usd_amt=750.0, fee_usd=-7.5, btc_amt=-0.025),
        ])
        try:
            _, txs = GeminiImporter().parse(path)
            assert len(txs) == 1
            tx = txs[0]
            assert tx['trans_type'] == 'Trade'
            assert tx['buy'] == pytest.approx(750.0)
            assert tx['buy_curr'] == 'USD'
            assert tx['sell'] == pytest.approx(0.025)   # abs of -0.025
            assert tx['sell_curr'] == 'BTC'
            assert tx['fee'] == pytest.approx(7.5)
            assert tx['fee_curr'] == 'USD'
        finally:
            os.unlink(path)

    def test_withdrawal_default_exchange(self):
        """DEBIT + BTC + Withdrawal spec, no withdraw_to → review comment."""
        dt = datetime(2024, 4, 1, 10, 0, 0)
        path = _make_xlsx([
            _xlsx_row(date=dt, typ="Debit", symbol="BTC",
                      spec="Withdrawal (BTC)",
                      btc_amt=-0.05, fee_btc=-0.0002, wd_dest="1A1zP1eP5QGefi2DMPTfTL5SLmv7DivfNa"),
        ])
        try:
            _, txs = GeminiImporter().parse(path, withdraw_to=None)
            assert len(txs) == 1
            tx = txs[0]
            assert tx['trans_type'] == 'Withdrawal'
            assert tx['exchange'] == 'Gemini-Withdrawal'
            assert tx['sell'] == pytest.approx(0.05)
            assert tx['sell_curr'] == 'BTC'
            assert tx['fee'] == pytest.approx(0.0002)
            assert tx['fee_curr'] == 'BTC'
            assert 'dest=1A1zP1eP5QGefi2DMPTfTL5SLmv7DivfNa' in tx['comment']
            assert 'Review' in tx['comment']
        finally:
            os.unlink(path)

    def test_withdrawal_with_withdraw_to(self):
        """Withdrawal with explicit withdraw_to → named exchange, no review."""
        dt = datetime(2024, 4, 2, 11, 0, 0)
        path = _make_xlsx([
            _xlsx_row(date=dt, typ="Debit", symbol="BTC",
                      spec="Withdrawal (BTC)",
                      btc_amt=-0.01, fee_btc=-0.0001, wd_dest="bc1qw508d6qejxtdg4y5r3zarvary0c5xw7kv8f3t4"),
        ])
        try:
            _, txs = GeminiImporter().parse(path, withdraw_to="ColdCard")
            tx = txs[0]
            assert tx['exchange'] == 'ColdCard'
            assert 'Review' not in tx['comment']
            assert 'dest=bc1qw508d6qejxtdg4y5r3zarvary0c5xw7kv8f3t4' in tx['comment']
        finally:
            os.unlink(path)

    def test_withdrawal_no_dest_address(self):
        """Withdrawal with empty Withdrawal Destination still works."""
        dt = datetime(2024, 4, 3, 12, 0, 0)
        path = _make_xlsx([
            _xlsx_row(date=dt, typ="Debit", symbol="BTC",
                      spec="Withdrawal (BTC)",
                      btc_amt=-0.005, fee_btc=0.0, wd_dest=""),
        ])
        try:
            _, txs = GeminiImporter().parse(path, withdraw_to=None)
            tx = txs[0]
            assert tx['trans_type'] == 'Withdrawal'
            assert 'dest=' not in tx['comment']
            assert 'Review' in tx['comment']
            assert tx['fee'] == 0.0
            assert tx['fee_curr'] == ''
        finally:
            os.unlink(path)

    def test_deposit_maps_correctly(self):
        """CREDIT + BTC → Deposit; no fee fields."""
        dt = datetime(2024, 5, 10, 8, 0, 0)
        path = _make_xlsx([
            _xlsx_row(date=dt, typ="Credit", symbol="BTC",
                      btc_amt=0.1),
        ])
        try:
            _, txs = GeminiImporter().parse(path)
            assert len(txs) == 1
            tx = txs[0]
            assert tx['trans_type'] == 'Deposit'
            assert tx['exchange'] == 'Gemini'
            assert tx['buy'] == pytest.approx(0.1)
            assert tx['buy_curr'] == 'BTC'
            assert tx['sell'] == 0.0
            assert tx['fee'] == 0.0
            assert tx['fee_curr'] == ''
        finally:
            os.unlink(path)

    def test_non_btc_rows_filtered(self):
        """Rows with Symbol != BTCUSD/BTC are silently skipped."""
        dt = datetime(2024, 6, 1, 10, 0, 0)
        # Add extra columns for GUSD so the file looks realistic
        headers = _XLSX_HEADERS + ["GUSD Amount GUSD"]
        path = _make_xlsx([
            # GUSD sell — not BTC
            [dt, "Sell", "GUSDUSD", "", 200.0, -2.0, 0.0, 0.0, "", 200.0],
            # BAT deposit — not BTC
            [dt, "Credit", "BAT", "", 0.0, 0.0, 0.0, 0.0, "", 0.0],
            # One real BTC buy mixed in
            [dt, "Buy", "BTCUSD", "", -300.0, -3.0, 0.01, 0.0, "", 0.0],
        ], headers=headers)
        try:
            _, txs = GeminiImporter().parse(path)
            assert len(txs) == 1
            assert txs[0]['trans_type'] == 'Trade'
            assert txs[0]['buy_curr'] == 'BTC'
        finally:
            os.unlink(path)

    def test_unknown_btc_type_skipped(self):
        """A BTC row with an unrecognised Type is silently dropped."""
        dt = datetime(2024, 7, 1, 10, 0, 0)
        path = _make_xlsx([
            _xlsx_row(date=dt, typ="Transfer", symbol="BTC", btc_amt=0.01),
        ])
        try:
            _, txs = GeminiImporter().parse(path)
            assert len(txs) == 0
        finally:
            os.unlink(path)

    def test_debit_without_withdrawal_in_spec_skipped(self):
        """DEBIT + BTC but spec doesn't contain 'Withdrawal' → skipped."""
        dt = datetime(2024, 7, 2, 10, 0, 0)
        path = _make_xlsx([
            _xlsx_row(date=dt, typ="Debit", symbol="BTC",
                      spec="Fee (BTC)", btc_amt=-0.0001),
        ])
        try:
            _, txs = GeminiImporter().parse(path)
            assert len(txs) == 0
        finally:
            os.unlink(path)

    def test_empty_xlsx_no_data_rows(self):
        """xlsx with headers only yields empty transaction list."""
        path = _make_xlsx([])
        try:
            colnames, txs = GeminiImporter().parse(path)
            assert len(txs) == 0
            assert len(colnames) == 11
        finally:
            os.unlink(path)

    def test_row_missing_date_skipped(self):
        """Row where Date is None is silently dropped."""
        path = _make_xlsx([
            _xlsx_row(date=None, typ="Buy", symbol="BTCUSD",
                      usd_amt=-100.0, btc_amt=0.003),
        ])
        try:
            _, txs = GeminiImporter().parse(path)
            assert len(txs) == 0
        finally:
            os.unlink(path)

    def test_multiple_buys_parsed(self):
        """Multiple BTC buys across different dates all parse correctly."""
        rows = [
            _xlsx_row(date=datetime(2024, 1, i+1, 10, 0, 0), typ="Buy",
                      symbol="BTCUSD", usd_amt=-100*(i+1),
                      fee_usd=-1.0, btc_amt=0.003*(i+1))
            for i in range(5)
        ]
        path = _make_xlsx(rows)
        try:
            _, txs = GeminiImporter().parse(path)
            assert len(txs) == 5
            # Verify amounts are distinct
            sells = [tx['sell'] for tx in txs]
            assert len(set(sells)) == 5
        finally:
            os.unlink(path)

    def test_fee_zero_yields_empty_fee_curr(self):
        """When fee is 0 the fee_curr field is empty string."""
        dt = datetime(2024, 8, 1, 10, 0, 0)
        path = _make_xlsx([
            _xlsx_row(date=dt, typ="Buy", symbol="BTCUSD",
                      usd_amt=-200.0, fee_usd=0.0, btc_amt=0.007),
        ])
        try:
            _, txs = GeminiImporter().parse(path)
            assert txs[0]['fee'] == 0.0
            assert txs[0]['fee_curr'] == ''
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# xlsx integration
# ---------------------------------------------------------------------------

class TestGeminiXlsxIntegration:
    """End-to-end xlsx import into in-memory SQLite."""

    def setup_method(self):
        clear_registry()
        register(GeminiImporter)

    def teardown_method(self):
        clear_registry()

    def _make_backend(self):
        return SqliteBackend(':memory:', auto_create_tables=True)

    def test_full_xlsx_import_workflow(self):
        """Parse a representative xlsx, validate, import; check all types land."""
        rows = [
            # Buy
            _xlsx_row(date=datetime(2024, 1, 5, 10, 0, 0), typ="Buy",
                      symbol="BTCUSD", usd_amt=-500.0, fee_usd=-5.0, btc_amt=0.016),
            # Sell
            _xlsx_row(date=datetime(2024, 1, 10, 14, 0, 0), typ="Sell",
                      symbol="BTCUSD", usd_amt=600.0, fee_usd=-6.0, btc_amt=-0.019),
            # Deposit
            _xlsx_row(date=datetime(2024, 1, 15, 8, 0, 0), typ="Credit",
                      symbol="BTC", btc_amt=0.05),
            # Withdrawal
            _xlsx_row(date=datetime(2024, 1, 20, 16, 0, 0), typ="Debit",
                      symbol="BTC", spec="Withdrawal (BTC)",
                      btc_amt=-0.03, fee_btc=-0.0002,
                      wd_dest="1BvBMSEYstWetqTFn5Au4m4GFg7xJaNVN2"),
        ]
        path = _make_xlsx(rows)
        try:
            parser = GeminiImporter()
            _, transactions = parser.parse(path, withdraw_to="Ledger")

            result = validate_batch(transactions)
            assert result.is_valid

            backend = self._make_backend()
            crypto = CryptoAccounts(backend=backend)
            import_result = crypto.import_transactions(transactions)

            assert import_result['imported'] == 4
            assert import_result['skipped'] == 0

            db_rows = backend.execute(
                "SELECT trans_type, COUNT(*) as cnt FROM ledger GROUP BY trans_type"
            )
            type_counts = {r['trans_type']: r['cnt'] for r in db_rows}
            assert type_counts.get('Trade', 0) == 2
            assert type_counts.get('Deposit', 0) == 1
            assert type_counts.get('Withdrawal', 0) == 1
        finally:
            os.unlink(path)

    def test_xlsx_withdrawal_dest_persists(self):
        """Withdrawal destination address survives the full import round-trip."""
        addr = "bc1qxy2kgdygjrsqtzq2nops68f327mx85345dxc8a"
        path = _make_xlsx([
            _xlsx_row(date=datetime(2024, 2, 1, 12, 0, 0), typ="Debit",
                      symbol="BTC", spec="Withdrawal (BTC)",
                      btc_amt=-0.01, fee_btc=-0.0001, wd_dest=addr),
        ])
        try:
            parser = GeminiImporter()
            _, transactions = parser.parse(path, withdraw_to="ColdCard")

            backend = self._make_backend()
            crypto = CryptoAccounts(backend=backend)
            crypto.import_transactions(transactions)

            rows = backend.execute(
                "SELECT comment FROM ledger WHERE trans_type = 'Withdrawal'"
            )
            assert len(rows) == 1
            assert f"dest={addr}" in rows[0]['comment']
        finally:
            os.unlink(path)
