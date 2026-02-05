"""Gemini transaction history parser — CSV and native xlsx.

Two export formats are supported:

1. CSV  (third-party gemini-exports tool)
   Columns: time, base-asset, quote-asset, type, price, quantity, total,
            fee, fee-currency, trade-id
   Filter: base-asset == BTC

2. xlsx  (native Gemini "Account History" download)
   30 columns with per-asset Amount/Fee/Balance columns.
   Key columns used: Date, Type, Symbol, Specification,
                     USD Amount USD, Fee (USD) USD,
                     BTC Amount BTC, Fee (BTC) BTC,
                     Withdrawal Destination
   Filter: Symbol in ('BTCUSD', 'BTC')
   Notes:
     - Buy USD amounts and fees are stored as negative; abs() is applied.
     - Type 'Credit' + Symbol 'BTC'  → Deposit
     - Type 'Debit'  + Symbol 'BTC'  → Withdrawal

Only BTC transactions are emitted by either path.
"""

from __future__ import annotations

import csv
from datetime import datetime as _datetime
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from typing import Any

from imports.base import BaseImporter
from imports.registry import register

# ---------------------------------------------------------------------------
# Shared output column list
# ---------------------------------------------------------------------------
_COLNAMES = [
    'trans_type', 'created_date', 'exchange', 'buy', 'buy_curr',
    'sell', 'sell_curr', 'fee', 'fee_curr', 'group', 'comment',
]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _parse_number(value: str) -> float:
    """Parse a numeric string, handling currency symbols and commas."""
    if not value:
        return 0.0
    value = value.strip().lstrip('$').replace(',', '')
    if not value:
        return 0.0
    try:
        return float(value)
    except ValueError:
        return 0.0


def _to_float(value: object) -> float:
    """Coerce an xlsx cell value (None / int / float / str) to float."""
    if value is None:
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _fmt_dt(value: object) -> str:
    """Format an xlsx datetime cell to 'YYYY-MM-DD HH:MM:SS'."""
    if isinstance(value, _datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    return str(value) if value else ''


# ---------------------------------------------------------------------------
# Parser class
# ---------------------------------------------------------------------------

@register
class GeminiImporter(BaseImporter):
    """Parser for Gemini transaction history (CSV or native xlsx)."""

    name = "Gemini"
    source_type = "exchange"
    file_patterns = ["*.csv", "*.xlsx"]
    description = "Gemini transaction history export (CSV or xlsx)"
    expected_columns = [
        "time", "base-asset", "quote-asset", "type", "price",
        "quantity", "total", "fee", "fee-currency", "trade-id",
    ]

    # ---- detection --------------------------------------------------------

    def detect(self, file_path: str) -> bool:
        """Return True for either the CSV or native xlsx Gemini format."""
        if file_path.lower().endswith('.xlsx'):
            return self._detect_xlsx(file_path)
        return self._detect_csv(file_path)

    def _detect_csv(self, file_path: str) -> bool:
        """CSV detection: look for hyphenated column names."""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                header = next(reader, None)
                if header is None:
                    return False
                header_lower = {col.lower().strip() for col in header}
                return (
                    'base-asset' in header_lower and
                    'quote-asset' in header_lower and
                    'trade-id' in header_lower
                )
        except (OSError, csv.Error):
            return False

    def _detect_xlsx(self, file_path: str) -> bool:
        """xlsx detection: look for Gemini native per-asset column names."""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path)
            ws = wb.active
            if ws is None or not ws.max_column:
                wb.close()
                return False
            headers = {
                str(ws.cell(row=1, column=c).value or '').lower().strip()
                for c in range(1, ws.max_column + 1)
            }
            wb.close()
            return 'btc amount btc' in headers and 'withdrawal destination' in headers
        except Exception:
            return False

    # ---- dispatch ---------------------------------------------------------

    def parse(
        self, file_path: str, withdraw_to: str | None = None
    ) -> tuple[list[str], list[dict[str, Any]]]:
        """Route to CSV or xlsx parser based on file extension."""
        if file_path.lower().endswith('.xlsx'):
            return self._parse_xlsx(file_path, withdraw_to)
        return self._parse_csv(file_path, withdraw_to)

    # ---- CSV path (unchanged from original) ------------------------------

    def _parse_csv(
        self, file_path: str, withdraw_to: str | None = None
    ) -> tuple[list[str], list[dict[str, Any]]]:
        transactions: list[dict[str, Any]] = []

        with open(file_path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            if not reader.fieldnames:
                return [], []

            col_map = {col.lower().strip(): col for col in reader.fieldnames}

            for row in reader:
                tx = self._parse_csv_row(row, col_map, withdraw_to)
                if tx is not None:
                    transactions.append(tx)

        return _COLNAMES, transactions

    def _parse_csv_row(
        self,
        row: dict[str, str],
        col_map: dict[str, str],
        withdraw_to: str | None,
    ) -> dict[str, Any] | None:
        """Parse a single CSV row.  Returns None for non-BTC or unknown types."""
        def get(col: str) -> str:
            original_col = col_map.get(col.lower(), col)
            return (row.get(original_col) or '').strip()

        base_asset = get('base-asset').upper()
        if base_asset != 'BTC':
            return None

        tx_type = get('type').upper()
        timestamp = get('time')
        quote_asset = get('quote-asset').upper() or 'USD'
        quantity = _parse_number(get('quantity'))
        total = _parse_number(get('total'))
        price = _parse_number(get('price'))
        fee = _parse_number(get('fee'))
        fee_currency = get('fee-currency').upper()

        if not timestamp or not tx_type:
            return None

        if tx_type == 'BUY':
            sell_amount = total if total > 0 else price * quantity
            return {
                'trans_type': 'Trade',
                'created_date': timestamp,
                'exchange': 'Gemini',
                'buy': quantity,
                'buy_curr': 'BTC',
                'sell': sell_amount,
                'sell_curr': quote_asset,
                'fee': fee,
                'fee_curr': fee_currency if fee > 0 else '',
                'group': '',
                'comment': '',
            }

        elif tx_type == 'SELL':
            buy_amount = total if total > 0 else price * quantity
            return {
                'trans_type': 'Trade',
                'created_date': timestamp,
                'exchange': 'Gemini',
                'buy': buy_amount,
                'buy_curr': quote_asset,
                'sell': quantity,
                'sell_curr': 'BTC',
                'fee': fee,
                'fee_curr': fee_currency if fee > 0 else '',
                'group': '',
                'comment': '',
            }

        elif tx_type == 'DEPOSIT':
            return {
                'trans_type': 'Deposit',
                'created_date': timestamp,
                'exchange': 'Gemini',
                'buy': quantity,
                'buy_curr': 'BTC',
                'sell': 0.0,
                'sell_curr': '',
                'fee': fee,
                'fee_curr': fee_currency if fee > 0 else '',
                'group': '',
                'comment': '',
            }

        elif tx_type == 'WITHDRAWAL':
            return {
                'trans_type': 'Withdrawal',
                'created_date': timestamp,
                'exchange': self._get_withdrawal_exchange(withdraw_to),
                'buy': 0.0,
                'buy_curr': '',
                'sell': quantity,
                'sell_curr': 'BTC',
                'fee': fee,
                'fee_curr': fee_currency if fee > 0 else '',
                'group': '',
                'comment': self._get_withdrawal_comment(withdraw_to),
            }

        elif tx_type in ('EARN', 'EARN INTEREST'):
            return {
                'trans_type': 'Interest Income',
                'created_date': timestamp,
                'exchange': 'Gemini',
                'buy': quantity,
                'buy_curr': 'BTC',
                'sell': 0.0,
                'sell_curr': '',
                'fee': 0.0,
                'fee_curr': '',
                'group': '',
                'comment': 'Gemini Earn',
            }

        return None

    # ---- xlsx path --------------------------------------------------------

    def _parse_xlsx(
        self, file_path: str, withdraw_to: str | None = None
    ) -> tuple[list[str], list[dict[str, Any]]]:
        """Parse the native Gemini 'Account History' xlsx export."""
        import openpyxl

        transactions: list[dict[str, Any]] = []
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        if ws is None or not ws.max_row or ws.max_row < 2:
            wb.close()
            return _COLNAMES, []

        # header name (lowered) → 1-based column index
        col_idx: dict[str, int] = {}
        for c in range(1, (ws.max_column or 0) + 1):
            name = ws.cell(row=1, column=c).value
            if name:
                col_idx[str(name).lower().strip()] = c

        for row_num in range(2, ws.max_row + 1):
            # Extract all values we need into a plain dict — avoids
            # closure-over-loop-variable issues and keeps _parse_xlsx_row pure.
            def _cell(h: str, rn: int = row_num) -> object:
                i = col_idx.get(h)
                return ws.cell(row=rn, column=i).value if i else None

            symbol = str(_cell('symbol') or '').upper().strip()
            if symbol not in ('BTCUSD', 'BTC'):
                continue

            data: dict[str, Any] = {
                'date': _cell('date'),
                'type': str(_cell('type') or '').upper().strip(),
                'specification': str(_cell('specification') or '').upper().strip(),
                'btc_amount': _to_float(_cell('btc amount btc')),
                'fee_btc': _to_float(_cell('fee (btc) btc')),
                'usd_amount': _to_float(_cell('usd amount usd')),
                'fee_usd': _to_float(_cell('fee (usd) usd')),
                'wd_dest': str(_cell('withdrawal destination') or '').strip(),
            }

            tx = self._parse_xlsx_row(data, symbol, withdraw_to)
            if tx is not None:
                transactions.append(tx)

        wb.close()
        return _COLNAMES, transactions

    def _parse_xlsx_row(
        self,
        data: dict[str, Any],
        symbol: str,
        withdraw_to: str | None,
    ) -> dict[str, Any] | None:
        """Map one native-xlsx row (pre-extracted dict) to a transaction dict."""
        typ = data['type']
        spec = data['specification']
        date_str = _fmt_dt(data['date'])

        if not date_str or not typ:
            return None

        btc_amt = data['btc_amount']
        fee_btc = data['fee_btc']
        usd_amt = data['usd_amount']
        fee_usd = data['fee_usd']
        wd_dest = data['wd_dest']

        # --- BTC Buy ---------------------------------------------------------
        if typ == 'BUY' and symbol == 'BTCUSD':
            return {
                'trans_type': 'Trade',
                'created_date': date_str,
                'exchange': 'Gemini',
                'buy': abs(btc_amt),
                'buy_curr': 'BTC',
                'sell': abs(usd_amt),       # stored as negative in xlsx
                'sell_curr': 'USD',
                'fee': abs(fee_usd),        # stored as negative in xlsx
                'fee_curr': 'USD' if fee_usd != 0 else '',
                'group': '',
                'comment': '',
            }

        # --- BTC Sell --------------------------------------------------------
        if typ == 'SELL' and symbol == 'BTCUSD':
            return {
                'trans_type': 'Trade',
                'created_date': date_str,
                'exchange': 'Gemini',
                'buy': abs(usd_amt),
                'buy_curr': 'USD',
                'sell': abs(btc_amt),       # stored as negative in xlsx
                'sell_curr': 'BTC',
                'fee': abs(fee_usd),
                'fee_curr': 'USD' if fee_usd != 0 else '',
                'group': '',
                'comment': '',
            }

        # --- BTC Withdrawal --------------------------------------------------
        if typ == 'DEBIT' and symbol == 'BTC' and 'WITHDRAWAL' in spec:
            parts: list[str] = []
            if wd_dest:
                parts.append(f"dest={wd_dest}")
            if not withdraw_to:
                parts.append("Review: Verify destination wallet")
            return {
                'trans_type': 'Withdrawal',
                'created_date': date_str,
                'exchange': self._get_withdrawal_exchange(withdraw_to),
                'buy': 0.0,
                'buy_curr': '',
                'sell': abs(btc_amt),       # stored as negative in xlsx
                'sell_curr': 'BTC',
                'fee': abs(fee_btc),
                'fee_curr': 'BTC' if fee_btc != 0 else '',
                'group': '',
                'comment': '; '.join(parts),
            }

        # --- BTC Deposit -----------------------------------------------------
        if typ == 'CREDIT' and symbol == 'BTC':
            return {
                'trans_type': 'Deposit',
                'created_date': date_str,
                'exchange': 'Gemini',
                'buy': abs(btc_amt),
                'buy_curr': 'BTC',
                'sell': 0.0,
                'sell_curr': '',
                'fee': 0.0,
                'fee_curr': '',
                'group': '',
                'comment': '',
            }

        # Unknown xlsx row type — skip
        return None
